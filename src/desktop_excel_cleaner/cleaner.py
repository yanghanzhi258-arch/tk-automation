"""Excel loading, data cleaning, and CSV export logic."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from .config import CSV_ENCODING, DEFAULT_OUTPUT_SUFFIX, SUPPORTED_EXCEL_EXTENSIONS
from .desktop import get_desktop_path

Row = dict[str, Any]


def normalize_column_name(name: Any, position: int) -> str:
    """Convert a raw Excel column name into a stable CSV column name."""

    if name is None:
        return f"column_{position}"

    normalized = " ".join(str(name).strip().split())
    return normalized or f"column_{position}"


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _normalize_value(value: Any) -> Any:
    return value.strip() if isinstance(value, str) else value


def clean_rows(
    headers: list[Any],
    rows: list[list[Any]],
    *,
    drop_duplicates: bool = True,
    sort_by: str | None = None,
) -> tuple[list[str], list[Row]]:
    """Clean spreadsheet rows using conservative, predictable rules."""

    normalized_headers = [
        normalize_column_name(header, index + 1) for index, header in enumerate(headers)
    ]

    non_empty_columns = [
        index
        for index, header in enumerate(normalized_headers)
        if not all(index >= len(row) or _is_empty(row[index]) for row in rows)
    ]
    cleaned_headers = [normalized_headers[index] for index in non_empty_columns]

    cleaned_rows: list[Row] = []
    seen: set[tuple[Any, ...]] = set()
    for row in rows:
        values = [
            _normalize_value(row[index]) if index < len(row) else None
            for index in non_empty_columns
        ]
        if all(_is_empty(value) for value in values):
            continue

        dedupe_key = tuple(values)
        if drop_duplicates and dedupe_key in seen:
            continue
        seen.add(dedupe_key)
        cleaned_rows.append(dict(zip(cleaned_headers, values, strict=True)))

    if sort_by:
        if sort_by not in cleaned_headers:
            available = ", ".join(cleaned_headers)
            raise ValueError(f"找不到排序列 '{sort_by}'。可用列：{available}")
        cleaned_rows.sort(key=lambda row: (row.get(sort_by) is None, row.get(sort_by)))

    return cleaned_headers, cleaned_rows


def build_output_path(input_path: Path, output_text: str | None = None) -> Path:
    """Build a desktop CSV output path from user input."""

    desktop = get_desktop_path()
    if output_text:
        output_path = Path(output_text).expanduser()
        if not output_path.is_absolute() and output_path.parent == Path("."):
            output_path = desktop / output_path
    else:
        output_path = desktop / f"{input_path.stem}{DEFAULT_OUTPUT_SUFFIX}"

    if output_path.suffix.lower() != ".csv":
        output_path = output_path.with_suffix(".csv")

    return output_path


def _read_excel_rows(input_path: Path, sheet_name: str | int) -> tuple[list[Any], list[list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "缺少依赖 openpyxl。请先运行：python -m pip install -r requirements.txt"
        ) from exc

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[sheet_name] if isinstance(sheet_name, int) else workbook[sheet_name]
        values = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not values:
        return [], []

    max_columns = max(len(row) for row in values)
    headers = list(values[0]) + [None] * (max_columns - len(values[0]))
    rows = [list(row) + [None] * (max_columns - len(row)) for row in values[1:]]
    return headers, rows


def convert_excel_to_csv(
    input_path: Path,
    *,
    sheet_name: str | int = 0,
    output_path: Path | None = None,
    drop_duplicates: bool = True,
    sort_by: str | None = None,
) -> Path:
    """Read an Excel file, clean it, and export the result as a CSV file."""

    if input_path.suffix.lower() not in SUPPORTED_EXCEL_EXTENSIONS:
        supported = ", ".join(SUPPORTED_EXCEL_EXTENSIONS)
        raise ValueError(f"不支持的文件类型：{input_path.suffix}。支持：{supported}")

    if not input_path.exists():
        raise FileNotFoundError(f"找不到 Excel 文件：{input_path}")

    headers, rows = _read_excel_rows(input_path, sheet_name)
    cleaned_headers, cleaned_rows = clean_rows(
        headers,
        rows,
        drop_duplicates=drop_duplicates,
        sort_by=sort_by,
    )

    target_path = output_path or build_output_path(input_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    with target_path.open("w", newline="", encoding=CSV_ENCODING) as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=cleaned_headers)
        writer.writeheader()
        writer.writerows(cleaned_rows)
    return target_path
